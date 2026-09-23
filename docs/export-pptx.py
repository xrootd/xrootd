#!/usr/bin/env python3
"""Export the XrdSecssh architecture deck to PPTX for Google Slides import."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree
import os

HERE = os.path.dirname(os.path.abspath(__file__))
XRD = os.path.join(HERE, "slides-assets", "xrootd-logo.png")
CERN = os.path.join(HERE, "slides-assets", "cern-logo.svg.png")
OUT = os.path.join(HERE, "XrdSecssh-architecture-slides.pptx")

BG = RGBColor(0x07, 0x07, 0x07)
INK = RGBColor(0xF7, 0xF4, 0xEF)
MUTED = RGBColor(0xC9, 0xBF, 0xB6)
GOLD = RGBColor(0xB8, 0x92, 0x3A)
PINK = RGBColor(0xE8, 0xA0, 0xBF)
PANEL = RGBColor(0x16, 0x14, 0x16)


def set_run(run, text, size=18, color=INK, bold=False, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = "Calibri"


def add_tb(slide, l, t, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, color, bold)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def outline_shape(shape, color, width_pt=1.25):
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def logos(slide, prs):
    if os.path.isfile(XRD):
        slide.shapes.add_picture(XRD, Inches(0.35), Inches(0.18), Inches(0.55), Inches(0.55))
    if os.path.isfile(CERN):
        slide.shapes.add_picture(CERN, Inches(12.45), Inches(0.18), Inches(0.52), Inches(0.52))
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.82), prs.slide_width, Pt(1.5))
    fill_shape(bar, GOLD)


def footer(slide, prs, n, total):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.28), prs.slide_width, Inches(0.22))
    fill_shape(bar, GOLD)
    add_tb(slide, 11.6, 7.18, 1.5, 0.28, f"{n} / {total}", size=11, color=GOLD, align=PP_ALIGN.RIGHT)


def kicker(slide, text):
    add_tb(slide, 0.55, 0.95, 12, 0.32, text.upper(), size=12, color=PINK, bold=True)


def title(slide, text, y=1.22, h=0.7, size=28):
    add_tb(slide, 0.55, y, 12.2, h, text, size=size, color=INK, bold=True)


def lead(slide, text, y=1.9, h=0.85, size=16):
    add_tb(slide, 0.55, y, 12.2, h, text, size=size, color=MUTED)


def card(slide, l, t, w, h, heading, body_lines, label=None, label_color=GOLD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    outline_shape(sh, GOLD, 1.0)
    y = t + 0.1
    if label:
        add_tb(slide, l + 0.12, y, w - 0.24, 0.22, label.upper(), size=10, color=label_color, bold=True)
        y += 0.22
    add_tb(slide, l + 0.12, y, w - 0.24, 0.32, heading, size=15, color=GOLD, bold=True)
    y += 0.34
    box = slide.shapes.add_textbox(Inches(l + 0.12), Inches(y), Inches(w - 0.24), Inches(h - (y - t) - 0.1))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in body_lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = 0
        p.space_after = Pt(4)
        run = p.add_run()
        prefix = "• " if not line.startswith("•") else ""
        set_run(run, prefix + line, 13, INK)


def bullets(slide, l, t, w, h, lines, size=16):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        run = p.add_run()
        set_run(run, "•  " + line, size, INK)


def table(slide, l, t, w, h, rows, col_w=None):
    n_rows, n_cols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_rows, n_cols, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = shp.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            is_hdr = r == 0
            set_run(run, val, 12 if not is_hdr else 11, GOLD if is_hdr else INK, bold=is_hdr)
            cell.text_frame.word_wrap = True
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            solid = etree.SubElement(tcPr, qn("a:solidFill"))
            srgb = etree.SubElement(solid, qn("a:srgbClr"))
            srgb.set("val", "161416" if r else "1E1810")
    return shp


def new_slide(prs, n, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    logos(slide, prs)
    footer(slide, prs, n, total)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    T = 16

    s = new_slide(prs, 1, T)
    add_tb(s, 0.9, 2.4, 11.5, 1.4, "XrdSecssh for XRootD", size=36, color=INK, bold=True)
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.85), Inches(1.6), Pt(3))
    fill_shape(rule, GOLD)
    add_tb(s, 0.9, 4.1, 11, 0.4, "Dr. Andreas-Joachim Peters - CERN IT-SD-PSS", size=20, color=PINK)
    add_tb(s, 0.9, 4.55, 11, 0.35, "XROOTD WORKSHOP LYON 2026", size=14, color=GOLD, bold=True)
    repo = add_tb(s, 0.9, 5.05, 11, 0.35, "github.com/cern-eos/xrootd - branch XrdSecssh", size=16, color=GOLD)
    repo.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://github.com/cern-eos/xrootd/tree/XrdSecssh"

    s = new_slide(prs, 2, T)
    kicker(s, "XRootD")
    title(s, "SSH-key login, over TLS only")
    lead(s, "sec.protocol ssh authenticates a client with an SSH public key or an OpenSSH user certificate. The server never sees a password. TLS authenticates the server and carries the handshake; the plugin maps a proven key to a local username.", y=1.95, h=0.95)
    card(s, 0.55, 3.1, 3.9, 3.5, "XrdSecssh", [
        "Who is this?",
        "Challenge/response over a key or user cert.",
        "Sets XrdSecEntity.name. No path decision.",
    ], "Authn", GOLD)
    card(s, 4.7, 3.1, 3.9, 3.5, "Keys or CA", [
        "keys-file maps raw keys to users.",
        "ca-keys-file validates OpenSSH user certificates.",
        "Either trust source is enough at init.",
    ], "Trust", GOLD)
    card(s, 8.85, 3.1, 3.9, 3.5, "Host in the signature", [
        "The client signs the hostname it connected to.",
        "A rogue redirect target cannot replay that signature onto another server.",
    ], "Bound", PINK)

    s = new_slide(prs, 3, T)
    kicker(s, "Talk")
    title(s, "Contents")
    bullets(s, 0.7, 2.05, 11, 4.6, [
        "1   Why SSH keys - and why not libssh",
        "2   Handshake: init, challenge, signed host",
        "3   Raw keys vs OpenSSH user certificates",
        "4   Principal mapping, policy, revocation",
        "5   Client keys, ssh-agent, configuration",
        "6   Hardening, tests and status",
    ], size=18)

    s = new_slide(prs, 4, T)
    kicker(s, "Design")
    title(s, "A key you already have, not another token")
    lead(s, "Unix and host protocols trust the connection. GSI and OIDC bring a PKI or an issuer. SSH keys are already on laptops and in authorized_keys. This plugin reuses that material on root:// over TLS.", y=1.9, h=0.75)
    table(s, 0.4, 2.75, 12.5, 3.7, [
        ["", "unix / host", "This plugin"],
        ["Who am I?", "Peer address or local uid", "Proven SSH key or user certificate"],
        ["Secret on the wire", "None / implicit", "Signature of nonce + host + fingerprint"],
        ["Server proof", "None in the sec handshake", "TLS server certificate (mandatory)"],
        ["Federation risk", "N/A", "Host binding stops a redirect-target relay"],
        ["HTTPS", "Separate", "Native root:// only; TLS is still required"],
    ], col_w=[2.5, 4.4, 5.6])
    add_tb(s, 0.55, 6.55, 12.2, 0.4, "Client authentication only. Do not disable TLS peer verification. Acc / ofs.authlib authorizes paths after login.", size=13, color=MUTED)

    s = new_slide(prs, 5, T)
    kicker(s, "Dependencies")
    title(s, "We use SSH key formats, not the SSH protocol")
    lead(s, "libssh and libssh2 implement RFC 4253: KEX, channels, SFTP, a second TCP stack. This plugin never opens an SSH session. It proves possession of a key over the TLS connection XRootD already has, using OpenSSL (already linked).", y=1.9, h=0.8)
    table(s, 0.4, 2.8, 12.5, 3.55, [
        ["", "libssh / libssh2", "What we actually need"],
        ["Transport", "SSH binary packet + KEX + MAC", "XRootD credentials on TLS"],
        ["Crypto", "Their own OpenSSL/gcrypt wrapper", "OpenSSL EVP - already in the tree"],
        ["Keys", "Full client identity + known_hosts", "Parse ssh-ed25519 / ssh-rsa blobs and OpenSSH certs"],
        ["Agent", "Whole agent client", "Two RPCs: list identities, sign (+ euid / mode checks)"],
        ["Policy", "OpenSSH server semantics", "Host binding, deny-users, fail-closed critical options"],
    ], col_w=[2.2, 4.7, 5.6])
    add_tb(s, 0.55, 6.5, 12.2, 0.5, "A second SSH library would add a packaging dependency, unused attack surface, and a handshake we would still have to replace. The wire format is xrdsec-ssh-v2, not SSH.", size=13, color=MUTED)

    s = new_slide(prs, 6, T)
    kicker(s, "Architecture")
    title(s, "Two round-trips, one per-connection challenge")
    lead(s, "Protocol version 1. Signed payload is xrdsec-ssh-v2 || string(host) || string(nonce) || string(fp). Challenge state lives on the protocol object, not in a global table.", y=1.9, h=0.7)
    card(s, 0.45, 2.7, 3.0, 1.4, "1  Init", ["user + key / cert blob"])
    card(s, 3.6, 2.7, 3.0, 1.4, "2  Challenge", ["nonce + fingerprint"])
    card(s, 6.75, 2.7, 3.0, 1.4, "3  Response", ["sig + connected host"])
    card(s, 9.9, 2.7, 2.95, 1.4, "4  Entity", ["name · prot=ssh"])
    card(s, 0.55, 4.3, 6.0, 2.1, "TLS is mandatory", [
        "needTLS() == true",
        "Non-TLS constructors refuse the protocol",
    ])
    card(s, 6.8, 4.3, 6.0, 2.1, "Verify", [
        "PendingChallenge: per object, single-use, -nonce-ttl",
        "host ∈ accepted names · signature · mapping",
    ])

    s = new_slide(prs, 7, T)
    kicker(s, "Trust")
    title(s, "Raw keys and user certificates")
    lead(s, "Init requires at least one usable trust source. Certificate-only sites may leave keys-file empty when -ca-keys-file loads.", y=1.9, h=0.65)
    card(s, 0.55, 2.7, 6.0, 3.8, "keys-file", [
        "alice ssh-ed25519 AAAA…",
        "or authorized_keys style ssh-ed25519 AAAA… alice@host",
        "Fingerprint must match a loaded key",
        "Requested user must match the mapped account",
        "Read once at plugin init",
    ], "Raw key", GOLD)
    card(s, 6.8, 2.7, 6.0, 3.8, "ca-keys-file", [
        "ssh-ed25519-cert-v01@openssh.com",
        "ssh-rsa-cert-v01@openssh.com",
        "CA signature, type=1, validity window",
        "Any critical option is a fail-closed reject",
        "Client must present the cert via ssh-agent",
    ], "User cert", PINK)
    add_tb(s, 0.55, 6.6, 12.2, 0.4, "V1: ssh-ed25519, ssh-rsa (≥2048). ECDSA and FIDO sk-* are not supported. RSA signatures must be rsa-sha2-256.", size=13, color=MUTED)

    s = new_slide(prs, 8, T)
    kicker(s, "Host binding")
    title(s, "A signature is only good for the host you reached")
    lead(s, "Without a host in the signed payload, a malicious data server (or a redirect you followed) can relay the client’s signature onto every other server that trusts the same keys.", y=1.9, h=0.75)
    card(s, 0.55, 2.8, 6.0, 3.7, "Signs the URL host", [
        "XrdSecProtocolsshObject gets the connected hostname",
        "Normalised: lower-case, no trailing dot, no [ipv6]",
        "Response carries that host next to the signature",
        "Challenge fingerprint must match the presented key",
    ], "Client", GOLD)
    card(s, 6.8, 2.8, 6.0, 3.7, "Accepts only its own names", [
        "FQDN, kernel hostname (long + short)",
        "localhost, 127.0.0.1, ::1",
        "Plus -hostnames a,b,… for aliases and VIPs",
        "Unknown host → “bound to a different server”",
    ], "Server", PINK)
    add_tb(s, 0.55, 6.6, 12.2, 0.4, "Not TLS channel binding. Closes the relay path while staying independent of the TLS stack.", size=13, color=MUTED)

    s = new_slide(prs, 9, T)
    kicker(s, "Certificates")
    title(s, "What must be true of a user cert")
    lead(s, "The CA is fully trusted. Restrict ca-keys-file. Empty principals are not a wildcard unless you opt in.", y=1.9, h=0.55)
    table(s, 0.45, 2.55, 12.4, 4.5, [
        ["Check", "Rule"],
        ["TLS", "Plugin constructors refuse non-TLS connections"],
        ["Type", "User certificate only (type=1)"],
        ["Validity", "valid_after / valid_before"],
        ["Critical options", "Any option (force-command, source-address, …) is rejected"],
        ["Principals", "Must contain the requested user if the list is non-empty"],
        ["Empty principals", "Rejected unless -allow-empty-principals"],
        ["RSA", "≥ 2048 bit subject key; CA sig label rsa-sha2-256"],
        ["Revocation", "Subject key, cert fingerprint, serial, key id"],
    ], col_w=[2.8, 9.6])

    s = new_slide(prs, 10, T)
    kicker(s, "Mapping")
    title(s, "From principal to local account")
    lead(s, "When mapping is enabled, the requested user is preferred if it is a listed, mappable principal. [root, alice] + request alice works.", y=1.9, h=0.65)
    card(s, 0.55, 2.7, 6.0, 3.8, "Direct and file", [
        "-principal-as-user - principal is a local name or uid",
        "-principal-map - /etc/xrootd/ssh_principals.map",
        "-principal-map-file <path>",
        "Direct mapping is tried first, then the file",
        "Map file is hot-reloaded (stat inode / mtime)",
    ], "Options", GOLD)
    card(s, 6.8, 2.7, 6.0, 3.8, "After every path", [
        "Username charset + length cap (64)",
        "-deny-users default root",
        "-deny-users none clears the list",
        "NSS lookups run outside the map mutex",
        "No global lock on the auth path",
    ], "Policy", PINK)

    s = new_slide(prs, 11, T)
    kicker(s, "Revocation")
    title(s, "A text list, hot-reloaded")
    lead(s, "Binary OpenSSH KRLs are not parsed. Use ssh-keygen -L for serial and key id, then list them.", y=1.9, h=0.55)
    card(s, 0.55, 2.6, 6.0, 4.0, "revoked-keys-file", [
        "ssh-ed25519 AAAA… - raw or subject key",
        "SHA256:base64fingerprint",
        "serial: 42",
        "id: alice-2026-01",
        "Same file-safety checks as the keys file",
    ])
    card(s, 6.8, 2.6, 6.0, 4.0, "Checked on init", [
        "Raw key fingerprint",
        "Certificate fingerprint",
        "Certificate subject-key fingerprint",
        "Serial and OpenSSH key id",
        "Client sees “key or certificate is revoked”",
    ])

    s = new_slide(prs, 12, T)
    kicker(s, "Client")
    title(s, "PEM, OpenSSH files, or ssh-agent")
    lead(s, "The first usable source wins. Certificate identities require the agent. Encrypted keys are refused - the client never prompts.", y=1.9, h=0.65)
    table(s, 0.45, 2.65, 12.4, 3.5, [
        ["#", "Source", "Notes"],
        ["1", "XRD_SSH_KEY_FILE", "OpenSSH native or PEM/PKCS8; alias XRD_SSH_PRIVATE_KEY_FILE"],
        ["2", "SSH_AUTH_SOCK", "Used if no key file, or when XRD_SSH_AGENT=1"],
        ["3", "~/.ssh/id_ed25519 then id_rsa", "Only if neither env is set and XRD_SSH_USER is unset"],
        ["-", "XRD_SSH_USER", "Else USER"],
        ["-", "XRD_SSH_AGENT_FINGERPRINT", "Pick one agent identity"],
    ], col_w=[1.0, 4.4, 7.0])
    add_tb(s, 0.55, 6.3, 12.2, 0.6, "Key files and agent sockets must be euid-owned and not group/other accessible. Keys are read from the checked fd (O_NOFOLLOW). Agent I/O times out after 10 s.", size=13, color=MUTED)

    s = new_slide(prs, 13, T)
    kicker(s, "Configuration")
    title(s, "sec.protocol ssh …")
    card(s, 0.55, 2.05, 6.0, 4.6, "xrootd.cf", [
        "xrootd.seclib libXrdSec.so",
        "xrootd.tls all",
        "sec.protocol ssh -keys-file /etc/xrootd/ssh_authorized_keys",
        "-ca-keys-file /etc/xrootd/ssh_ca_keys",
        "-revoked-keys-file /etc/xrootd/ssh_revoked_keys",
        "-principal-as-user -principal-map",
        "-hostnames data1.example.org,vip",
        "-deny-users root,daemon",
    ])
    card(s, 6.8, 2.05, 6.0, 4.6, "Defaults and ranges", [
        "keys-file /etc/xrootd/ssh_authorized_keys",
        "-maxsz 8192 (1 … 524288)",
        "-nonce-ttl 30 s (1 … 600)",
        "keys / CA files: init only",
        "principal-map and revoked: hot-reload",
        "Files: regular, euid owner, not group/other writable, O_NOFOLLOW, ≤ 10 MB",
    ])

    s = new_slide(prs, 14, T)
    kicker(s, "Hardening")
    title(s, "What we refuse to do on the hot path")
    table(s, 0.4, 2.05, 12.5, 4.7, [
        ["Risk", "Mitigation"],
        ["Signature relay via redirect", "Host in the signed payload; accepted-name set + -hostnames"],
        ["Global challenge table DoS", "Per-object PendingChallenge; dies with the connection"],
        ["Empty-principal wildcard / root", "Reject empty principals by default; -deny-users root"],
        ["Mapping leak to the client", "Generic “SSH authentication failed.”; detail in the server log"],
        ["Small RSA / SHA-1 labels", "≥ 2048 bit; rsa-sha2-256 only"],
        ["Key-file TOCTOU / symlink", "O_NOFOLLOW, fstat, read the checked fd"],
        ["Wedged ssh-agent", "10 s send/receive timeout"],
        ["Parser DoS", "64 KiB wire fields; line / base64 / cred size caps"],
    ], col_w=[4.6, 7.9])
    add_tb(s, 0.55, 6.85, 12.2, 0.3, "Debug (-debug / XrdSecDEBUG=1) redacts fingerprints and omits usernames and socket paths.", size=12, color=MUTED)

    s = new_slide(prs, 15, T)
    kicker(s, "Tests")
    title(s, "Unit coverage in xrdsecssh-unit-tests")
    lead(s, "The test TU includes the protocol object. CMake links the vendored GTest::gtest targets (same as master).", y=1.9, h=0.65)
    card(s, 0.55, 2.7, 3.9, 3.6, "Wire / crypto", [
        "ed25519 / RSA sign-verify",
        "Host binding and aliases",
        "Trailing bytes, maxsz",
        "Per-object challenge lifetime",
        "OpenSSH-format and encrypted keys",
    ])
    card(s, 4.7, 2.7, 3.9, 3.6, "Policy", [
        "Empty principals default / opt-in",
        "deny-users on raw and cert",
        "Requested-principal preference",
        "Generic client errors",
        "Small RSA rejected everywhere",
    ])
    card(s, 8.85, 2.7, 3.9, 3.6, "Ops", [
        "Safe file reader / symlink",
        "Cert type, expiry, critical opts",
        "Revocation parse + hot-reload",
        "Init option ranges",
        "Cert-only init without keys-file",
    ])
    add_tb(s, 0.55, 6.45, 12.2, 0.4, "68 tests.  ctest -R xrdsecssh-unit-tests", size=13, color=MUTED)

    s = new_slide(prs, 16, T)
    kicker(s, "Outlook")
    title(s, "Status")
    card(s, 0.55, 2.05, 3.9, 3.7, "Done", [
        "sec.protocol ssh + TLS gate",
        "Raw keys and OpenSSH user certs",
        "Host-bound challenge (v1 wire)",
        "Principal map, deny-users, revoke",
        "OpenSSH private-key files + agent",
        "68 unit tests, packaging",
    ], "Done", GOLD)
    card(s, 4.7, 2.05, 3.9, 3.7, "Not done", [
        "TLS channel binding / exporter",
        "ECDSA / FIDO sk-*",
        "Binary OpenSSH KRL files",
        "HTTPS / XrdHttp handshake",
        "Encrypted private keys (use agent)",
    ], "Not done", PINK)
    card(s, 8.85, 2.05, 3.9, 3.7, "Next", [
        "Propose the plugin for xrootd/master",
        "Always set -hostnames for aliases",
        "Keep CAs tight; prefer mapping + deny-users",
        "Docs: src/XrdSecssh/README.md",
    ], "Next", PINK)
    add_tb(s, 0.55, 5.95, 12.2, 0.85, "XrdSecssh authenticates the client with an SSH key or user certificate on TLS root://. The server is authenticated by TLS. The signed host name stops a federation relay. Acc authorizes paths afterwards.", size=15, color=MUTED)

    prs.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
